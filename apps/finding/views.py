import json
import csv
from django.shortcuts import render
from apps.tables.models import Finding
from django.contrib.contenttypes.models import ContentType
from django.http import HttpResponse
from django.shortcuts import render, redirect
from django.core.paginator import Paginator
from django.contrib.auth.decorators import login_required
from apps.tables.utils import (
    software_filter, same_key_filter, common_date_filter,
    common_float_filter, common_integer_filter, 
    common_count_filter, common_unique_filter, get_model_fields
)
from apps.tables.models import UserFilter, PageItems, HideShowFilter, IntRangeFilter, FloatRangeFilter, ModelChoices, DateRangeFilter, SelectedRows
from django.http import HttpResponse
from django.urls import reverse
from django.http import JsonResponse
from datetime import datetime, timedelta
from django.utils import timezone
from loader.models import InstantUpload
from django.contrib.contenttypes.models import ContentType
from dateutil.parser import isoparse
from apps.common.models import SavedFilter, FieldType
from openpyxl import Workbook
from django.views import View

# Create your views here.

def get_user_id(request):
    if request.user.is_authenticated:
        return request.user.pk
    else:
        return -1

def create_finding_filter(request, finding_id):
    finding = Finding.objects.get(id=finding_id)
    content_type = ContentType.objects.get(app_label=finding.content_type.app_label, model=finding.content_type.model)
    model = content_type.model_class()

    field_type_map = {
        'date': getattr(model, 'date_fields_to_convert', []),
        'int': getattr(model, 'integer_fields', []),
        'float': getattr(model, 'float_fields', [])
    }

    def get_field_type(value):
        for f_type, fields in field_type_map.items():
            if value in fields:
                return getattr(FieldType, f_type.upper())
        return FieldType.TEXT

    def get_boolean_field(name, index):
        return request.POST.get(f'{name}_{index}') == 'true'

    existing_count_filter = SavedFilter.objects.filter(
        userID=request.user.id, parent=ModelChoices.FINDING, is_count=True, finding_id=finding.id
    ).first()
    for key, value in request.POST.items():
        if key.startswith('field_name_') and value:
            index = key.split('_')[-1]
            field_name = value
            field_type = get_field_type(value)

            if field_type == FieldType.TEXT:
                value_start = request.POST.get(f'value_start_{index}') or ''
                value_end = request.POST.get(f'value_end_{index}') or ''
            else:
                value_start = request.POST.get(f'value_start_{index}') or None
                value_end = request.POST.get(f'value_end_{index}') or None

            is_null = get_boolean_field('is_null', index)
            is_not_null = get_boolean_field('is_not_null', index)
            is_unique = get_boolean_field('is_unique', index)
            is_count = get_boolean_field('is_count', index)

            filter_id = request.POST.get(f'filter_id_{index}')
            if filter_id in [None, '', 'null']:
                filter_id = None
            
            if is_count:
                if existing_count_filter and (not filter_id or str(existing_count_filter.id) != filter_id):
                    is_count = False

            filter_data = {
                'userID': request.user.id,
                'parent': ModelChoices.FINDING,
                'field_name': field_name,
                'field_type': field_type,
                'value_start': value_start,
                'value_end': value_end,
                'is_null': is_null,
                'is_not_null': is_not_null,
                'is_unique': is_unique,
                'is_count': is_count,
                'finding_id': finding.id
            }

            if filter_id:
                SavedFilter.objects.filter(pk=filter_id).update(**filter_data)
                saved_filter = SavedFilter.objects.get(pk=filter_id)
            else:
                saved_filter = SavedFilter.objects.create(**filter_data)
            
            if not finding.saved_filters.filter(pk=saved_filter.pk).exists():
                finding.saved_filters.add(saved_filter)
    
    return redirect(request.META.get('HTTP_REFERER'))

def create_finding_page_items(request, finding_id):
    finding = Finding.objects.get(id=finding_id)
    if request.method == 'POST':
        items = request.POST.get('items')
        page_items, created = PageItems.objects.update_or_create(
            userID=get_user_id(request),
            parent=ModelChoices.FINDING,
            finding_id=finding.id,
            defaults={'items_per_page':items}
        )
        finding.page_items = page_items
        finding.save()

        return redirect(reverse('list_of_devices'))

def create_finding_hide_show_filter(request, finding_id):
    if request.method == "POST":
        data_list = json.loads(request.body)
        finding = Finding.objects.get(id=finding_id)

        for data in data_list:
            hide_show_filter, created = HideShowFilter.objects.update_or_create(
                userID=get_user_id(request),
                parent=ModelChoices.FINDING,
                key=data.get('key'),
                finding_id=finding.id,
                defaults={'value': data.get('value')}
            )
            finding.hide_show_filters.add(hide_show_filter)

        response_data = {'message': 'Filters updated successfully'}
        return JsonResponse(response_data)

    return JsonResponse({'error': 'Invalid request'}, status=400)


def delete_finding_filter(request):
    if request.method == 'POST':
        filter_id = request.POST.get('filter_id')
        try:
            # Delete the filter from the database
            filter_to_delete = SavedFilter.objects.get(id=filter_id, userID=request.user.id)
            filter_to_delete.delete()
            return JsonResponse({'success': True})
        except SavedFilter.DoesNotExist:
            return JsonResponse({'success': False, 'error': 'Filter not found'})
    return JsonResponse({'success': False, 'error': 'Invalid request method'})

def delete_finding_date_range_filter(request, id):
    filter_instance = DateRangeFilter.objects.get(id=id, userID=get_user_id(request), parent=ModelChoices.FINDING)
    filter_instance.delete()

    return redirect(request.META.get('HTTP_REFERER'))

def delete_finding_int_range_filter(request, id):
    filter_instance = IntRangeFilter.objects.get(id=id, userID=get_user_id(request), parent=ModelChoices.FINDING)
    filter_instance.delete()

    return redirect(request.META.get('HTTP_REFERER'))

def delete_finding_float_range_filter(request, id):
    filter_instance = FloatRangeFilter.objects.get(id=id, userID=get_user_id(request), parent=ModelChoices.FINDING)
    filter_instance.delete()

    return redirect(request.META.get('HTTP_REFERER'))



@login_required(login_url='/users/signin/')
def add_to_finding(request, model_name, model_choice):
    if request.method == 'POST':
        content_type = ContentType.objects.get(model=model_name.lower())
        img_loader_id = request.POST.get("img_loader_id")
        tab_id = request.POST.get("tab_id")
        
        user_filter_qs = UserFilter.objects.filter(
            userID=get_user_id(request), parent=getattr(ModelChoices, model_choice),
            **({"img_loader_id": img_loader_id} if img_loader_id else {}),
            **({"tab_id": tab_id} if tab_id else {}),
        )
        date_range_filter_qs = DateRangeFilter.objects.filter(
            userID=get_user_id(request), parent=getattr(ModelChoices, model_choice),
            **({"img_loader_id": img_loader_id} if img_loader_id else {}),
            **({"tab_id": tab_id} if tab_id else {}),
        )

        int_range_filter_qs = IntRangeFilter.objects.filter(
            userID=get_user_id(request), parent=getattr(ModelChoices, model_choice),
            **({"img_loader_id": img_loader_id} if img_loader_id else {}),
            **({"tab_id": tab_id} if tab_id else {}),
        )

        float_range_filter_qs = FloatRangeFilter.objects.filter(
            userID=get_user_id(request), parent=getattr(ModelChoices, model_choice),
            **({"img_loader_id": img_loader_id} if img_loader_id else {}),
            **({"tab_id": tab_id} if tab_id else {}),
        )

        saved_filter_qs = SavedFilter.objects.filter(
            userID=get_user_id(request), parent=getattr(ModelChoices, model_choice),
            **({"img_loader_id": img_loader_id} if img_loader_id else {}),
            **({"tab_id": tab_id} if tab_id else {}),
        )

        hide_show_filter_qs = HideShowFilter.objects.filter(
            userID=get_user_id(request), parent=getattr(ModelChoices, model_choice),
            **({"img_loader_id": img_loader_id} if img_loader_id else {}),
            **({"tab_id": tab_id} if tab_id else {}),
        )

        page_items_qs = PageItems.objects.filter(
            userID=get_user_id(request), parent=getattr(ModelChoices, model_choice),
            **({"img_loader_id": img_loader_id} if img_loader_id else {}),
            **({"tab_id": tab_id} if tab_id else {}),
        ).last()
        
        fav_name = request.POST.get('fav_name')

        finding = Finding.objects.create(
            user=request.user,
            name=fav_name,
            content_type=content_type,
            model_choices=model_choice,
            search=request.POST.get('search_fields'),
            pre_columns=request.POST.get('pre_columns'),
            pre_filters=request.POST.get('pre_filters'),
            order_by=request.POST.get('order_by'),
            selected_rows=request.POST.get('selected_rows', ''),
            snapshot=request.POST.get('snapshot') if request.POST.get('snapshot') else 'latest',
            created_by=request.user.email if request.user.email else request.user.username,
            created_at=timezone.now(),
            updated_at=timezone.now()
        )

        if request.POST.get('query_snapshot'):
            finding.query_snapshot = request.POST.get('query_snapshot')
        
        if request.POST.get('is_dynamic_query') == 'true':
            finding.is_dynamic_query = True

        if page_items_qs:
            page_items = PageItems.objects.create(
                userID=page_items_qs.userID,
                parent=ModelChoices.FINDING,
                items_per_page=page_items_qs.items_per_page,
                finding_id=finding.pk
            )

            finding.page_items = page_items

        for user_filter in user_filter_qs:
            user_filters = UserFilter.objects.create(
                userID=user_filter.userID,
                parent=ModelChoices.FINDING,
                key=user_filter.key,
                value=user_filter.value,
                finding_id=finding.pk
            )

            finding.user_filters.add(user_filters)
    
        for date_range_filter in date_range_filter_qs:
            date_range_filters = DateRangeFilter.objects.create(
                userID=date_range_filter.userID,
                parent=ModelChoices.FINDING,
                from_date=date_range_filter.from_date,
                to_date=date_range_filter.to_date,
                key=date_range_filter.key,
                finding_id=finding.pk
            )

            finding.date_range_filters.add(date_range_filters)
        
        for int_range_filter in int_range_filter_qs:
            int_range_filters = IntRangeFilter.objects.create(
                userID=int_range_filter.userID,
                parent=ModelChoices.FINDING,
                from_number=int_range_filter.from_number,
                to_number=int_range_filter.to_number,
                key=int_range_filter.key,
                finding_id=finding.pk
            )

            finding.int_range_filters.add(int_range_filters)
        
        for float_range_filter in float_range_filter_qs:
            float_range_filters = FloatRangeFilter.objects.create(
                userID=float_range_filter.userID,
                parent=ModelChoices.FINDING,
                from_float_number=float_range_filter.from_float_number,
                to_float_number=float_range_filter.to_float_number,
                key=float_range_filter.key,
                finding_id=finding.pk
            )

            finding.float_range_filters.add(float_range_filters)
        

        for saved_filter in saved_filter_qs:
            saved_filters = SavedFilter.objects.create(
                userID=saved_filter.userID,
                parent=ModelChoices.FINDING,
                field_name=saved_filter.field_name,
                field_type=saved_filter.field_type,
                value_start=saved_filter.value_start,
                value_end=saved_filter.value_end,
                is_null=saved_filter.is_null,
                is_not_null=saved_filter.is_not_null,
                is_unique=saved_filter.is_unique,
                is_count=saved_filter.is_count,
                finding_id=finding.pk
            )
            finding.saved_filters.add(saved_filters)
        
        for hide_show_filter in hide_show_filter_qs:
            hide_show_filters = HideShowFilter.objects.create(
                userID=hide_show_filter.userID,
                parent=ModelChoices.FINDING,
                key=hide_show_filter.key,
                value=hide_show_filter.value,
                finding_id=finding.pk
            )

            finding.hide_show_filters.add(hide_show_filters)
        
        finding.save()

    return redirect(reverse('update_finding', args=[finding.id]))
    # return redirect(request.META.get('HTTP_REFERER'))


@login_required(login_url='/users/signin/')
def add_finding_from_finding(request, finding_id):
    finding = Finding.objects.get(id=finding_id)

    if request.method == 'POST':
        if request.POST.get('search_fields'):
            finding.search = request.POST.get('search_fields')
        
        if request.POST.get('order_by') :
            finding.order_by = request.POST.get('order_by') 

        if request.POST.get('snapshot'):
            finding.snapshot = request.POST.get('snapshot') 

        if finding.page_items:
            page_items = PageItems.objects.get(
                userID=finding.page_items.userID,
                parent=ModelChoices.FINDING,
                pk=finding.page_items.pk
            )

            finding.page_items = page_items

        if finding.user_filters.all():
            for user_filter in finding.user_filters.all():
                user_filters = UserFilter.objects.get(
                    userID=user_filter.userID,
                    parent=ModelChoices.FINDING,
                    pk=user_filter.pk
                )

                finding.user_filters.add(user_filters)

        if finding.date_range_filters.all():
            for date_range_filter in finding.date_range_filters.all():
                date_range_filters = DateRangeFilter.objects.get(
                    userID=date_range_filter.userID,
                    parent=ModelChoices.FINDING,
                    pk=date_range_filter.pk
                )

                finding.date_range_filters.add(date_range_filters)

        if finding.int_range_filters.all():
            for int_range_filter in finding.int_range_filters.all():
                int_range_filters = IntRangeFilter.objects.get(
                    userID=int_range_filter.userID,
                    parent=ModelChoices.FINDING,
                    pk=int_range_filter.pk
                )

                finding.int_range_filters.add(int_range_filters)
        

        if finding.float_range_filters.all():
            for float_range_filter in finding.float_range_filters.all():
                float_range_filters = FloatRangeFilter.objects.get(
                    userID=float_range_filter.userID,
                    parent=ModelChoices.FINDING,
                    pk=float_range_filter.pk
                )

                finding.float_range_filters.add(float_range_filters)
        
        if finding.hide_show_filters.all():
            for hide_show_filter in finding.hide_show_filters.all():
                hide_show_filters = HideShowFilter.objects.get(
                    userID=hide_show_filter.userID,
                    parent=ModelChoices.FINDING,
                    pk=hide_show_filter.pk
                )

                finding.hide_show_filters.add(hide_show_filters)
        
        finding.updated_at = timezone.now()
        finding.save()

    return redirect(request.META.get('HTTP_REFERER'))

def finding_details(request, id):
    finding = Finding.objects.get(id=id)
    content_type = ContentType.objects.get(app_label=finding.content_type.app_label, model=finding.content_type.model)
    db_field_names = [field.name for field in content_type.model_class()._meta.get_fields() if not field.is_relation and not '_Q_' in field.name]
    finding_model = content_type.model_class()

    # Snapshot
    latest_snapshot = ''
    summary = None
    snapshot_filter = {}

    if not finding.is_dynamic_query:
        try:
            content_type = ContentType.objects.get(model=finding_model.__name__.lower())
            snapshots = InstantUpload.objects.filter(content_type=content_type).order_by('-created_at')
            snapshot = request.GET.get('snapshot')

            if snapshot and not snapshot == 'all':
                summary = InstantUpload.objects.get(id=snapshot)
                snapshot_filter['loader_instance'] = summary.pk
            
            elif snapshot and snapshot == 'all':
                snapshot_filter= {}
            
            elif finding.snapshot == 'latest':
                latest_snapshot = snapshots.latest('created_at')
                snapshot_filter['loader_instance'] = latest_snapshot.pk
            
            elif finding.snapshot and not finding.snapshot == 'all':
                summary = InstantUpload.objects.get(id=int(finding.snapshot))
                snapshot_filter['loader_instance'] = int(finding.snapshot)
            
            elif finding.snapshot and finding.snapshot == 'all':
                snapshot_filter= {}
            

        except:
            pass
    
    else:
        snapshots = finding_model.objects.exclude(snapshot=None).values('snapshot').distinct()
    

    items = 25
    if finding.page_items:
        items = finding.page_items.items_per_page

    filter_string = {}
    combined_q_objects, user_unique_filter, user_query_conditions, user_count_filters = same_key_filter(finding.saved_filters.filter(field_type=FieldType.TEXT), return_count_filters=True)

    pre_filters = {}
    if finding.pre_filters:
        pre_filters = eval(finding.pre_filters)
    

    if finding.is_dynamic_query:
        query_snapshot = finding.query_snapshot
        if request.GET.get('query_snapshot'):
            query_snapshot = request.GET.get('query_snapshot')
        else:
            query_snapshot = finding.query_snapshot
        
        if query_snapshot and query_snapshot != 'all':
            parsed_datetime = isoparse(query_snapshot)
            snapshot_filter['snapshot'] = parsed_datetime
    
    # for date range
    date_string, date_unique_filter, date_query_conditions, date_count_filters = common_date_filter(finding.saved_filters.filter(field_type=FieldType.DATE), return_count_filters=True)
    filter_string.update(date_string)
    
    # for integer range
    int_string, int_unique_filter, int_query_conditions, int_count_filters = common_integer_filter(finding.saved_filters.filter(field_type=FieldType.INT), return_count_filters=True)
    filter_string.update(int_string)

    # for float range
    float_string, float_unique_filter, float_query_conditions, float_count_filters = common_float_filter(finding.saved_filters.filter(field_type=FieldType.FLOAT), return_count_filters=True)
    filter_string.update(float_string)
    
    if request.GET.get('search'):
        finding.search = request.GET.get('search')
        finding.updated_at = timezone.now()
        finding.save()
    
    if request.GET.get('order_by'):
        finding.order_by = request.GET.get('order_by')
        finding.updated_at = timezone.now()
        finding.save()

    base_queryset = finding_model.objects.filter(combined_q_objects).filter(**filter_string).filter(**snapshot_filter).filter(**pre_filters)
    if finding.selected_rows:
        base_queryset = base_queryset.filter(pk__in=eval(finding.selected_rows))

    order_by = request.GET.get('order_by', 'pk')
    queryset = base_queryset
    if hasattr(finding_model, 'parent'):
        queryset = queryset.filter(parent=None)

    if user_query_conditions:
        queryset = queryset.filter(user_query_conditions)
    if date_query_conditions:
        queryset = queryset.filter(date_query_conditions)
    if int_query_conditions:
        queryset = queryset.filter(int_query_conditions)
    if float_query_conditions:
        queryset = queryset.filter(float_query_conditions)

    if user_count_filters:
        queryset = common_count_filter(user_count_filters, base_queryset, queryset, db_field_names)
    elif date_count_filters:
        queryset = common_count_filter(date_count_filters, base_queryset, queryset, db_field_names)
    elif int_count_filters:
        queryset = common_count_filter(int_count_filters, base_queryset, queryset, db_field_names)
    elif float_count_filters:
        queryset = common_count_filter(float_count_filters, base_queryset, queryset, db_field_names)
    else:
        if order_by in ['count', '-count']:
            order_by = 'pk'

    queryset = queryset.order_by(order_by)

    table_name = f"{finding.content_type.app_label}_{finding.content_type.model}"
    if user_unique_filter:
        queryset = common_unique_filter(request, user_unique_filter, queryset, snapshot_filter, table_name)
    if date_unique_filter:
        queryset = common_unique_filter(request, date_unique_filter, queryset, snapshot_filter, table_name)
    if int_unique_filter:
        queryset = common_unique_filter(request, int_unique_filter, queryset, snapshot_filter, table_name)
    if float_unique_filter:
        queryset = common_unique_filter(request, float_unique_filter, queryset, snapshot_filter, table_name)

    finding_list = software_filter(request, queryset, db_field_names)

    page = request.GET.get('page', 1)
    paginator = Paginator(finding_list, items)
    
    finding_items = paginator.page(page)

    pre_column = finding.pre_columns
    richtext_fields = finding.richtext_fields if finding.richtext_fields else []

    fields = get_model_fields(f'{finding_model.__name__}', pre_column)
    saved_filters = list(SavedFilter.objects.filter(
        userID=request.user.id, parent=ModelChoices.FINDING, finding_id=finding.id
    ).values())
    for filter in saved_filters:
        if 'created_at' in filter:
            filter['created_at'] = filter['created_at'].isoformat()
    saved_filters_json = json.dumps(saved_filters)

    context = {
        'finding': finding,
        'db_field_names': db_field_names,
        'finding_items': finding_items,
        'pre_column': pre_column,
        'richtext_fields': richtext_fields,
        'total_items': base_queryset.count(),
        'date_picker_fields': finding_model.date_fields_to_convert if finding_model.date_fields_to_convert else [],
        'int_fields': finding_model.integer_fields if finding_model.integer_fields else [],
        'float_fields': finding_model.float_fields if finding_model.float_fields else [],
        'snapshots': snapshots,
        'selected_snapshot': summary,
        'latest_snapshot': latest_snapshot,
        'is_dynamic_query': finding.is_dynamic_query,
        'query_snapshot': query_snapshot if finding.is_dynamic_query else '',
        'fields': fields,
        'saved_filters': saved_filters_json,
        'unique_filter': list(set(user_unique_filter + date_unique_filter + int_unique_filter + float_unique_filter)),
        'table_name': finding.content_type.model,
    }
    return render(request, 'apps/finding/test.html', context)


def remove_finding(request, id):
    if request.method == 'POST':
        finding = Finding.objects.get(id=id)
        finding.delete()

        return redirect('/index')


def add_finding_description(request, finding_id):
    finding = Finding.objects.get(id=finding_id)
    if request.method == 'POST':
        description = request.POST.get('description')
        finding.description = description
        finding.updated_at = timezone.now()
        finding.save()
        
    return redirect(request.META.get('HTTP_REFERER')) 



def export_finding_csv_view(request, id):
    finding = Finding.objects.get(id=id)
    content_type = ContentType.objects.get(app_label=finding.content_type.app_label, model=finding.content_type.model)
    db_field_names = [field.name for field in content_type.model_class()._meta.get_fields() if not field.is_relation and not '_Q_' in field.name]
    finding_model = content_type.model_class()

    # Snapshot
    latest_snapshot = ''
    summary = None
    snapshot_filter = {}
    try:
        content_type = ContentType.objects.get(model=finding_model.__name__.lower())
        snapshots = InstantUpload.objects.filter(content_type=content_type).order_by('-created_at')
        snapshot = request.GET.get('snapshot')

        if snapshot and not snapshot == 'all':
            summary = InstantUpload.objects.get(id=snapshot)
            snapshot_filter['loader_instance'] = summary.pk
        
        elif snapshot and snapshot == 'all':
            snapshot_filter= {}
        
        elif finding.snapshot == 'latest':
            latest_snapshot = snapshots.latest('created_at')
            snapshot_filter['loader_instance'] = latest_snapshot.pk
        
        elif finding.snapshot and not finding.snapshot == 'all':
            summary = InstantUpload.objects.get(id=int(finding.snapshot))
            snapshot_filter['loader_instance'] = int(finding.snapshot)
        
        elif finding.snapshot and finding.snapshot == 'all':
            snapshot_filter= {}
        

    except:
        pass

    fields = []
    pre_column = finding.pre_columns
    show_fields = finding.hide_show_filters.all().filter(value=False)
    for field in show_fields:
        if field.key not in pre_column:
            fields.append(field.key)

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="{finding_model._meta.model_name}.csv"'

    
    filter_string = {}
    combined_q_objects, user_unique_filter, user_query_conditions, user_count_filters = same_key_filter(finding.saved_filters.filter(field_type=FieldType.TEXT), return_count_filters=True)

    pre_filters = {}
    if finding.pre_filters:
        pre_filters = eval(finding.pre_filters)

    # for date range
    date_string, date_unique_filter, date_query_conditions, date_count_filters = common_date_filter(finding.saved_filters.filter(field_type=FieldType.DATE), return_count_filters=True)
    filter_string.update(date_string)
    
    # for integer range
    int_string, int_unique_filter, int_query_conditions, int_count_filters = common_integer_filter(finding.saved_filters.filter(field_type=FieldType.INT), return_count_filters=True)
    filter_string.update(int_string)

    # for float range
    float_string, float_unique_filter, float_query_conditions, float_count_filters = common_float_filter(finding.saved_filters.filter(field_type=FieldType.FLOAT), return_count_filters=True)
    filter_string.update(float_string)

    base_queryset = finding_model.objects.filter(**filter_string).filter(**snapshot_filter).filter(combined_q_objects).filter(**pre_filters)
    order_by = request.GET.get('order_by', 'pk')
    queryset = base_queryset

    if user_query_conditions:
        queryset = queryset.filter(user_query_conditions)
    if date_query_conditions:
        queryset = queryset.filter(date_query_conditions)
    if int_query_conditions:
        queryset = queryset.filter(int_query_conditions)
    if float_query_conditions:
        queryset = queryset.filter(float_query_conditions)

    if user_count_filters:
        queryset = common_count_filter(user_count_filters, base_queryset, queryset, fields)
    elif date_count_filters:
        queryset = common_count_filter(date_count_filters, base_queryset, queryset, fields)
    elif int_count_filters:
        queryset = common_count_filter(int_count_filters, base_queryset, queryset, fields)
    elif float_count_filters:
        queryset = common_count_filter(float_count_filters, base_queryset, queryset, fields)
    else:
        if order_by in ['count', '-count']:
            order_by = 'pk'
        if 'count' in fields:
            fields.remove('count')

    queryset = queryset.order_by(order_by)

    table_name = f"{finding.content_type.app_label}_{finding.content_type.model}"
    if user_unique_filter:
        queryset = common_unique_filter(request, user_unique_filter, queryset, snapshot_filter, table_name)
    if date_unique_filter:
        queryset = common_unique_filter(request, date_unique_filter, queryset, snapshot_filter, table_name)
    if int_unique_filter:
        queryset = common_unique_filter(request, int_unique_filter, queryset, snapshot_filter, table_name)
    if float_unique_filter:
        queryset = common_unique_filter(request, float_unique_filter, queryset, snapshot_filter, table_name)

    if finding.selected_rows:
        queryset = queryset.filter(ID__in=eval(finding.selected_rows))

    finding_list = software_filter(request, queryset, db_field_names)

    writer = csv.writer(response)
    writer.writerow(fields) 

    for fav in finding_list:
        row_data = []
        for field in fields:
            if field in finding_model.date_fields_to_convert:
                unix_timestamp = getattr(fav, field)
                if unix_timestamp:
                    date_time = datetime.utcfromtimestamp(unix_timestamp)
                    formatted_date = date_time.strftime('%b %d %Y %I:%M%p')
                    row_data.append(formatted_date)
                else: 
                        row_data.append("")
            else:
                row_data.append(getattr(fav, field))
        writer.writerow(row_data)

    return response

class ExportFindingExcelView(View):
    def get_excel_workbook(self, request, id):
        """Returns just the workbook for easier integration"""
        finding = Finding.objects.get(id=id)
        content_type = ContentType.objects.get(app_label=finding.content_type.app_label, model=finding.content_type.model)
        finding_model = content_type.model_class()
        db_field_names = [field.name for field in finding_model._meta.get_fields() 
                         if not field.is_relation and '_Q_' not in field.name]

        # Selected rows handling
        selected_rows = SelectedRows.objects.filter(
            model=finding.content_type.model, 
            model_choice=ModelChoices.TAB,
            finding_id=finding.id
        ).values_list('rows', flat=True)
        selected_rows = [int(item) for row in selected_rows for item in row.split(',')]

        # Snapshot handling
        latest_snapshot = ''
        summary = None
        snapshot_filter = {}
        try:
            snapshots = InstantUpload.objects.filter(content_type=content_type).order_by('-created_at')
            snapshot = request.GET.get('snapshot')

            if snapshot and snapshot != 'all':
                summary = InstantUpload.objects.get(id=snapshot)
                snapshot_filter['loader_instance'] = summary.pk
            elif snapshot == 'all':
                snapshot_filter = {}
            elif finding.snapshot == 'latest':
                latest_snapshot = snapshots.latest('created_at')
                snapshot_filter['loader_instance'] = latest_snapshot.pk
            elif finding.snapshot and finding.snapshot != 'all':
                summary = InstantUpload.objects.get(id=int(finding.snapshot))
                snapshot_filter['loader_instance'] = int(finding.snapshot)
            elif finding.snapshot == 'all':
                snapshot_filter = {}
        except:
            pass

        # Field filtering
        fields = []
        pre_column = finding.pre_columns
        richtext_fields = finding.richtext_fields or []
        show_fields = finding.hide_show_filters.filter(value=False)
        show_fields_keys = {field.key for field in show_fields}

        for field in db_field_names:
            if field not in pre_column and field in show_fields_keys:
                fields.append(field)

        # Filter handling
        filter_string = {}
        combined_q_objects, user_unique_filter, user_query_conditions, user_count_filters = same_key_filter(
            finding.saved_filters.filter(field_type=FieldType.TEXT), return_count_filters=True)

        pre_filters = eval(finding.pre_filters) if finding.pre_filters else {}

        date_string, date_unique_filter, date_query_conditions, date_count_filters = common_date_filter(
            finding.saved_filters.filter(field_type=FieldType.DATE), return_count_filters=True)
        filter_string.update(date_string)

        int_string, int_unique_filter, int_query_conditions, int_count_filters = common_integer_filter(
            finding.saved_filters.filter(field_type=FieldType.INT), return_count_filters=True)
        filter_string.update(int_string)

        float_string, float_unique_filter, float_query_conditions, float_count_filters = common_float_filter(
            finding.saved_filters.filter(field_type=FieldType.FLOAT), return_count_filters=True)
        filter_string.update(float_string)

        # Query construction
        base_queryset = finding_model.objects.filter(**filter_string).filter(**snapshot_filter)\
                                         .filter(combined_q_objects).filter(**pre_filters)
        order_by = request.GET.get('order_by', 'pk')
        queryset = base_queryset

        if user_query_conditions:
            queryset = queryset.filter(user_query_conditions)
        if date_query_conditions:
            queryset = queryset.filter(date_query_conditions)
        if int_query_conditions:
            queryset = queryset.filter(int_query_conditions)
        if float_query_conditions:
            queryset = queryset.filter(float_query_conditions)

        if user_count_filters:
            queryset = common_count_filter(user_count_filters, base_queryset, queryset, fields)
        elif date_count_filters:
            queryset = common_count_filter(date_count_filters, base_queryset, queryset, fields)
        elif int_count_filters:
            queryset = common_count_filter(int_count_filters, base_queryset, queryset, fields)
        elif float_count_filters:
            queryset = common_count_filter(float_count_filters, base_queryset, queryset, fields)
        else:
            if order_by in ['count', '-count']:
                order_by = 'pk'
            if 'count' in fields:
                fields.remove('count')

        queryset = queryset.order_by(order_by)

        # Unique filters
        table_name = f"{finding.content_type.app_label}_{finding.content_type.model}"
        if user_unique_filter:
            queryset = common_unique_filter(request, user_unique_filter, queryset, snapshot_filter, table_name)
        if date_unique_filter:
            queryset = common_unique_filter(request, date_unique_filter, queryset, snapshot_filter, table_name)
        if int_unique_filter:
            queryset = common_unique_filter(request, int_unique_filter, queryset, snapshot_filter, table_name)
        if float_unique_filter:
            queryset = common_unique_filter(request, float_unique_filter, queryset, snapshot_filter, table_name)

        finding_list = software_filter(request, queryset, db_field_names)

        if 'selected_rows' in request.GET:
            finding_list = finding_list.filter(pk__in=selected_rows)

        # Workbook creation
        wb = Workbook()
        ws = wb.active
        ws.append(fields)

        for fav in finding_list:
            row_data = []
            for field in fields:
                if field in finding_model.date_fields_to_convert:
                    unix_timestamp = getattr(fav, field)
                    if unix_timestamp:
                        epoch = datetime(1970, 1, 1, tzinfo=timezone.utc)
                        date_time = epoch + timedelta(seconds=unix_timestamp)
                        formatted_date = date_time.strftime('%b %d %Y %I:%M%p')
                        row_data.append(formatted_date)
                    else:
                        row_data.append("")
                elif field in richtext_fields:
                    value = getattr(fav, field, None)
                    row_data.append(value.html if hasattr(value, 'html') else "")
                else:
                    row_data.append(getattr(fav, field))
            ws.append(row_data)

        return wb

    def get(self, request, id):
        """Original endpoint that returns HTTP response"""
        wb = self.get_excel_workbook(request, id)
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        finding = Finding.objects.get(id=id)
        content_type = ContentType.objects.get(app_label=finding.content_type.app_label, model=finding.content_type.model)
        finding_model = content_type.model_class()
        response['Content-Disposition'] = f'attachment; filename="{finding_model._meta.model_name}.xlsx"'
        wb.save(response)
        return response
