import json
import csv
from django.contrib.contenttypes.models import ContentType
from apps.common.models import SavedFilter, FieldType
from django.http import HttpResponse
from django.shortcuts import render, redirect
from django.core.paginator import Paginator
from django.contrib.auth.decorators import login_required
from apps.tables.utils import (
    software_filter, same_key_filter, common_date_filter,
    common_float_filter, common_integer_filter, common_count_filter, 
    common_unique_filter, get_model_fields
)
from apps.tables.models import (
    PageItems, HideShowFilter,  ModelChoices, 
    Tab, ActionStatus, SelectedRows, TabNotes
)
from django.http import HttpResponse
from django.urls import reverse
from django.http import JsonResponse
from datetime import datetime, timedelta
from django.utils import timezone
from loader.models import InstantUpload
from django.contrib.contenttypes.models import ContentType
from dateutil.parser import isoparse
from apps.common.models import SavedFilter
from django.contrib import messages
from django.db.models import F, Case, When, IntegerField
from openpyxl import Workbook
from django.views import View
from home.models import ColumnOrder
from apps.tables.forms import TabNotesForm


def get_user_id(request):
    if request.user.is_authenticated:
        return request.user.pk
    else:
        return -1

def create_tab_filter(request, tab_id):
    tab = Tab.objects.get(id=tab_id)
    content_type = ContentType.objects.get(app_label=tab.content_type.app_label, model=tab.content_type.model)
    model = content_type.model_class()

    field_type_map = {
        'date': getattr(model, 'date_fields_to_convert', []),
        'int': getattr(model, 'integer_fields', []),
        'float': getattr(model, 'float_fields', [])
    }

    def get_field_type(value):
        for f_type, fields in field_type_map.items():
            if value in fields:
                return getattr(FieldType, f_type.upper())
        return FieldType.TEXT

    def get_boolean_field(name, index):
        return request.POST.get(f'{name}_{index}') == 'true'
    
    ipe_path = reverse('tab_details', args=[tab.id])
    last_ipe = IPE.objects.filter(path=ipe_path).last()
    base_description = last_ipe.description if last_ipe else ""

    existing_count_filter = SavedFilter.objects.filter(
        userID=request.user.id, parent=ModelChoices.TAB, is_count=True, tab_id=tab.id
    ).first()

    filters_description = []

    for key, value in request.POST.items():
        if key.startswith('field_name_') and value:
            index = key.split('_')[-1]
            field_name = value
            field_type = get_field_type(value)

            if field_type == FieldType.TEXT:
                value_start = request.POST.get(f'value_start_{index}') or ''
                value_end = request.POST.get(f'value_end_{index}') or ''
            else:
                value_start = request.POST.get(f'value_start_{index}') or None
                value_end = request.POST.get(f'value_end_{index}') or None

            is_null = get_boolean_field('is_null', index)
            is_not_null = get_boolean_field('is_not_null', index)
            is_not = get_boolean_field('is_not', index)
            is_unique = get_boolean_field('is_unique', index)
            is_count = get_boolean_field('is_count', index)

            filter_id = request.POST.get(f'filter_id_{index}')
            if filter_id in [None, '', 'null']:
                filter_id = None
            
            if is_count:
                if existing_count_filter and (not filter_id or str(existing_count_filter.id) != filter_id):
                    is_count = False

            filter_data = {
                'userID': request.user.id,
                'parent': ModelChoices.TAB,
                'field_name': field_name,
                'field_type': field_type,
                'value_start': value_start,
                'value_end': value_end,
                'is_null': is_null,
                'is_not_null': is_not_null,
                'is_not': is_not,
                'is_unique': is_unique,
                'is_count': is_count,
                'tab_id': tab.id
            }

            if filter_id:
                SavedFilter.objects.filter(pk=filter_id).update(**filter_data)
                saved_filter = SavedFilter.objects.get(pk=filter_id)
            else:
                saved_filter = SavedFilter.objects.create(**filter_data)
            
            if not tab.saved_filters.filter(pk=saved_filter.pk).exists():
                tab.saved_filters.add(saved_filter)
            
            filters_description.append(
                f"Field: {field_name}, Start: {value_start}, End: {value_end}, "
                f"Null: {is_null}, Not Null: {is_not_null}, Not: {saved_filter.is_not}, Unique: {is_unique}, Count: {is_count}"
            )
    
    user = request.user
    timestamp_utc = timezone.now()
    timestamp_local = timezone.localtime(timestamp_utc)
    timestamp_str = timestamp_local.strftime('%Y-%m-%d %H:%M:%S')

    new_description = (
        f"{base_description}\n" if base_description else ""
    ) + f'Filters applied by: {user.get_full_name() or user.username}\n' \
        f'Applied on: {timestamp_str}\n' \
        f'Filters:\n' + "\n".join(filters_description)

    IPE.objects.create(
        userID=user,
        path=ipe_path,
        tab_id=tab.id,
        description=new_description
    )
    
    return redirect(request.META.get('HTTP_REFERER'))


def create_tab_page_items(request, tab_id):
    tab = Tab.objects.get(id=tab_id)
    if request.method == 'POST':
        items = request.POST.get('items')
        page_items, created = PageItems.objects.update_or_create(
            userID=get_user_id(request),
            parent=ModelChoices.TAB,
            tab_id=tab.id,
            defaults={'items_per_page':items}
        )
        tab.page_items = page_items
        tab.save()

        return redirect(request.META.get('HTTP_REFERER'))

def create_tab_hide_show_filter(request, tab_id):
    if request.method == "POST":
        data_list = json.loads(request.body)
        tab = Tab.objects.get(id=tab_id)

        for data in data_list:
            hide_show_filter, created = HideShowFilter.objects.update_or_create(
                userID=get_user_id(request),
                parent=ModelChoices.TAB,
                key=data.get('key'),
                tab_id=tab.id,
                defaults={'value': data.get('value')}
            )
            tab.hide_show_filters.add(hide_show_filter)

        response_data = {'message': 'Filters updated successfully'}
        return JsonResponse(response_data)

    return JsonResponse({'error': 'Invalid request'}, status=400)

def delete_tab_filter(request):
    if request.method == 'POST':
        filter_id = request.POST.get('filter_id')
        try:
            filter_to_delete = SavedFilter.objects.get(id=filter_id, userID=request.user.id)
             # ipe_path = reverse('tab_details', args=[filter_to_delete.tab_id])
            # last_ipe = IPE.objects.filter(path=ipe_path).last()
            # base_description = last_ipe.description if last_ipe else ""

            # user = request.user
            # timestamp_utc = timezone.now()
            # timestamp_local = timezone.localtime(timestamp_utc)
            # timestamp_str = timestamp_local.strftime('%Y-%m-%d %H:%M:%S')

            # filters_description = (
            #     f"Field: {filter_to_delete.field_name}, Start: {filter_to_delete.value_start}, End: {filter_to_delete.value_end}, "
            #     f"Null: {filter_to_delete.is_null}, Not Null: {filter_to_delete.is_not_null}, Unique: {filter_to_delete.is_unique}, Count: {filter_to_delete.is_count}"
            # )

            # new_description = (
            #     f"{base_description}\n" if base_description else ""
            # ) + f'Filters applied by: {user.get_full_name() or user.username}\n' \
            #     f'Applied on: {timestamp_str}\n' \
            #     f'Filters:\n{filters_description}'

            # IPE.objects.create(
            #     userID=user,
            #     path=ipe_path,
            #     tab_id=filter_to_delete.tab_id,
            #     description=new_description
            # )
            filter_to_delete.delete()
            return JsonResponse({'success': True})
        except SavedFilter.DoesNotExist:
            return JsonResponse({'success': False, 'error': 'Filter not found'})
    return JsonResponse({'success': False, 'error': 'Invalid request method'})


import io
from home.models import IPE
from django.utils.timezone import now
from urllib.parse import urlparse
from django.urls import reverse

@login_required(login_url='/users/signin/')
def add_to_tab(request, model_name, model_choice):
    if request.method == 'POST':
        content_type = ContentType.objects.get(model=model_name.lower())
        img_loader_id = request.POST.get("img_loader_id")
        saved_filter_qs = SavedFilter.objects.filter(
            userID=get_user_id(request), parent=getattr(ModelChoices, model_choice),
            **({"img_loader_id": img_loader_id} if img_loader_id else {})
        )

        hide_show_filter_qs = HideShowFilter.objects.filter(
            userID=get_user_id(request), parent=getattr(ModelChoices, model_choice),
            **({"img_loader_id": img_loader_id} if img_loader_id else {})
        )

        page_items_qs = PageItems.objects.filter(
            userID=get_user_id(request), parent=getattr(ModelChoices, model_choice),
            **({"img_loader_id": img_loader_id} if img_loader_id else {})
        ).last()
        
        fav_name = request.POST.get('fav_name')
        base_view = request.POST.get('base_view')

        if Tab.objects.filter(name=fav_name, base_view=base_view).exists():
            Tab.objects.filter(name=fav_name, base_view=base_view).delete()

        tab = Tab.objects.create(
            user=request.user,
            name=fav_name,
            content_type=content_type,
            model_choices=model_choice,
            search=request.POST.get('search_fields'),
            pre_columns=request.POST.get('pre_columns'),
            pre_filters=request.POST.get('pre_filters'),
            base_view=base_view,
            sidebar_parent=request.POST.get('parent'),
            richtext_fields=request.POST.get('richtext_fields'),
            order_by=request.POST.get('order_by'),
            snapshot=request.POST.get('snapshot') if request.POST.get('snapshot') else 'latest',
            created_at=timezone.now(),
            updated_at=timezone.now()
        )

        if request.POST.get('query_snapshot'):
            tab.query_snapshot = request.POST.get('query_snapshot')
        
        if request.POST.get('is_dynamic_query') == 'true':
            tab.is_dynamic_query = True

        if page_items_qs:
            page_items = PageItems.objects.create(
                userID=page_items_qs.userID,
                parent=ModelChoices.TAB,
                items_per_page=page_items_qs.items_per_page,
                tab_id=tab.pk
            )

            tab.page_items = page_items
        
        filters_description = []
        if saved_filter_qs.exists():
            for saved_filter in saved_filter_qs:
                saved_filters = SavedFilter.objects.create(
                    userID=saved_filter.userID,
                    parent=ModelChoices.TAB,
                    field_name=saved_filter.field_name,
                    field_type=saved_filter.field_type,
                    value_start=saved_filter.value_start,
                    value_end=saved_filter.value_end,
                    is_null=saved_filter.is_null,
                    is_not_null=saved_filter.is_not_null,
                    is_not=saved_filter.is_not,
                    is_unique=saved_filter.is_unique,
                    is_count=saved_filter.is_count,
                    tab_id=tab.pk
                )
                tab.saved_filters.add(saved_filters)

                filters_description.append(
                    f"Field: {saved_filter.field_name}, Start: {saved_filter.value_start}, End: {saved_filter.value_end}, "
                    f"Null: {saved_filter.is_null}, Not Null: {saved_filter.is_not_null}, Not: {saved_filter.is_not}, Unique: {saved_filter.is_unique}, Count: {saved_filter.is_count}"
                )

        for hide_show_filter in hide_show_filter_qs:
            hide_show_filters = HideShowFilter.objects.create(
                userID=hide_show_filter.userID,
                parent=ModelChoices.TAB,
                key=hide_show_filter.key,
                value=hide_show_filter.value,
                tab_id=tab.pk
            )

            tab.hide_show_filters.add(hide_show_filters)
        
        tab.save()


        user = request.user
        timestamp_utc = timezone.now()
        timestamp_local = timezone.localtime(timestamp_utc)
        timestamp_str = timestamp_local.strftime('%Y-%m-%d %H:%M:%S')

        base_ipe = IPE.objects.filter(path=urlparse(request.META.get('HTTP_REFERER', request.path)).path).last()
        base_description = base_ipe.description if base_ipe else ""
        new_description = (
            f"{base_description}\n" if base_description else ""
        ) + f'\nTab Created by: {user.get_full_name() or user.username}\n' \
            f'Tab Created on: {timestamp_str}\n' \
            f'Tab name: {tab.name}\n' \
            f'Base Filters:\n' + "\n".join(filters_description)
                        
        IPE.objects.create(
            userID=user,
            path=reverse('tab_details', args=[tab.id]),
            tab_id=tab.id,
            description=new_description               
        )

        # create tab note
        TabNotes.objects.get_or_create(tab=tab)

    return redirect(request.META.get('HTTP_REFERER'))


@login_required(login_url='/users/signin/')
def add_tab_from_tab(request, tab_id):
    tab = Tab.objects.get(id=tab_id)

    if request.method == 'POST':
        if request.POST.get('search_fields'):
            tab.search = request.POST.get('search_fields')
        
        if request.POST.get('order_by') :
            tab.order_by = request.POST.get('order_by') 

        if request.POST.get('snapshot'):
            tab.snapshot = request.POST.get('snapshot') 

        if tab.page_items:
            page_items = PageItems.objects.get(
                userID=tab.page_items.userID,
                parent=ModelChoices.TAB,
                pk=tab.page_items.pk
            )

            tab.page_items = page_items
        
        if tab.hide_show_filters.all():
            for hide_show_filter in tab.hide_show_filters.all():
                hide_show_filters = HideShowFilter.objects.get(
                    userID=hide_show_filter.userID,
                    parent=ModelChoices.TAB,
                    pk=hide_show_filter.pk
                )

                tab.hide_show_filters.add(hide_show_filters)
        
        tab.updated_at = timezone.now()
        tab.save()

    return redirect(request.META.get('HTTP_REFERER'))

def row_count_to_ipe(request, tab, count):
    if tab.saved_filters.exists():
        ipe_path = reverse('tab_details', args=[str(tab.id)])
        ipe = IPE.objects.filter(path=ipe_path, tab_id=tab.id).last()

        if ipe:
            ipe.description = (ipe.description or '') + f'\nRows after filter: {count}\n'
            ipe.save()

def tab_details(request, id):
    tab = Tab.objects.get(id=id)
    tab_notes, _ = TabNotes.objects.get_or_create(tab=tab)
    content_type = ContentType.objects.get(app_label=tab.content_type.app_label, model=tab.content_type.model)
    db_field_names = [field.name for field in content_type.model_class()._meta.get_fields() if not field.is_relation]
    tab_model = content_type.model_class()

    try:
        user_order = ColumnOrder.objects.get(user=request.user, table_name=f'{tab_model.__name__}', tab_id=f"{tab.id}")
        column_names = [col['key'] for col in user_order.column_order if col['key'] is not None]
        ordered_fields = column_names
    except ColumnOrder.DoesNotExist:
        ordered_fields = db_field_names

    field_dict = {field.key: field for field in tab.hide_show_filters.all()}
    field_names = [field_dict[key] for key in ordered_fields if key in field_dict]

    selected_rows = SelectedRows.objects.filter(
        model=f'{tab.content_type.model}', 
        model_choice=ModelChoices.TAB,
        tab_id=tab.id
    ).values_list('rows', flat=True)
    selected_rows = [int(item) for row in selected_rows for item in row.split(',')]

    # Snapshot
    latest_snapshot = ''
    summary = None
    snapshot_filter = {}

    if not tab.is_dynamic_query:
        try:
            content_type = ContentType.objects.get(model=tab_model.__name__.lower())
            snapshots = InstantUpload.objects.filter(content_type=content_type).order_by('-created_at')
            snapshot = request.GET.get('snapshot')

            if snapshot and not snapshot == 'all':
                summary = InstantUpload.objects.get(id=snapshot)
                snapshot_filter['loader_instance'] = summary.pk
            
            elif snapshot and snapshot == 'all':
                snapshot_filter= {}
            
            elif tab.snapshot == 'latest':
                latest_snapshot = snapshots.latest('created_at')
                snapshot_filter['loader_instance'] = latest_snapshot.pk
            
            elif tab.snapshot and not tab.snapshot == 'all':
                summary = InstantUpload.objects.get(id=int(tab.snapshot))
                snapshot_filter['loader_instance'] = int(tab.snapshot)
            
            elif tab.snapshot and tab.snapshot == 'all':
                snapshot_filter= {}
            
        except:
            pass
    
    else:
        snapshots = tab_model.objects.exclude(snapshot=None).values('snapshot').distinct()
    
    items = 25
    if tab.page_items:
        items = tab.page_items.items_per_page

    filter_string = {}
    combined_q_objects, user_unique_filter, user_query_conditions, user_count_filters = same_key_filter(tab.saved_filters.filter(field_type=FieldType.TEXT), return_count_filters=True)

    pre_filters = {}
    if tab.pre_filters:
        pre_filters = eval(tab.pre_filters)
    
    if tab.is_dynamic_query:
        query_snapshot = tab.query_snapshot
        if request.GET.get('query_snapshot'):
            query_snapshot = request.GET.get('query_snapshot')
        else:
            query_snapshot = tab.query_snapshot
        

        if query_snapshot and query_snapshot != 'all':
            parsed_datetime = isoparse(query_snapshot)
            snapshot_filter['snapshot'] = parsed_datetime
    
    # for date range
    date_string, date_unique_filter, date_query_conditions, date_count_filters = common_date_filter(tab.saved_filters.filter(field_type=FieldType.DATE), return_count_filters=True)
    filter_string.update(date_string)
    
    # for integer range
    int_string, int_unique_filter, int_query_conditions, int_count_filters = common_integer_filter(tab.saved_filters.filter(field_type=FieldType.INT), return_count_filters=True)
    filter_string.update(int_string)

    # for float range
    float_string, float_unique_filter, float_query_conditions, float_count_filters = common_float_filter(tab.saved_filters.filter(field_type=FieldType.FLOAT), return_count_filters=True)
    filter_string.update(float_string)

    if request.GET.get('search'):
        tab.search = request.GET.get('search')
        tab.updated_at = timezone.now()
        tab.save()
    
    if request.GET.get('order_by'):
        tab.order_by = request.GET.get('order_by')
        tab.updated_at = timezone.now()
        tab.save()

    base_queryset = tab_model.objects.filter(combined_q_objects).filter(**filter_string).filter(**snapshot_filter).filter(**pre_filters)
    order_by = request.GET.get('order_by', 'pk')
    queryset = base_queryset
    if hasattr(tab_model, 'parent'):
        queryset = queryset.filter(parent=None)
        try:
            queryset = queryset.filter(action_status=ActionStatus.IS_ACTIVE)
        except:
            pass
    
    if user_query_conditions:
        queryset = queryset.filter(user_query_conditions)
    if date_query_conditions:
        queryset = queryset.filter(date_query_conditions)
    if int_query_conditions:
        queryset = queryset.filter(int_query_conditions)
    if float_query_conditions:
        queryset = queryset.filter(float_query_conditions)

    if user_count_filters:
        queryset = common_count_filter(user_count_filters, base_queryset, queryset, ordered_fields)
    elif date_count_filters:
        queryset = common_count_filter(date_count_filters, base_queryset, queryset, ordered_fields)
    elif int_count_filters:
        queryset = common_count_filter(int_count_filters, base_queryset, queryset, ordered_fields)
    elif float_count_filters:
        queryset = common_count_filter(float_count_filters, base_queryset, queryset, ordered_fields)
    else:
        if order_by in ['count', '-count']:
            order_by = 'pk'

    queryset = queryset.order_by(order_by)

    table_name = f"{tab.content_type.app_label}_{tab.content_type.model}"
    if user_unique_filter:
        queryset = common_unique_filter(request, user_unique_filter, queryset, snapshot_filter, table_name)
    if date_unique_filter:
        queryset = common_unique_filter(request, date_unique_filter, queryset, snapshot_filter, table_name)
    if int_unique_filter:
        queryset = common_unique_filter(request, int_unique_filter, queryset, snapshot_filter, table_name)
    if float_unique_filter:
        queryset = common_unique_filter(request, float_unique_filter, queryset, snapshot_filter, table_name)

    if selected_rows:
        queryset = queryset.annotate(
            order_priority=Case(
                *[When(pk=row_id, then=0) for row_id in selected_rows],
                default=1,
                output_field=IntegerField(),
            )
        ).order_by('order_priority')

    tab_list = software_filter(request, queryset, ordered_fields)
    row_count_to_ipe(request, tab, tab_list.count())

    page = request.GET.get('page', 1)
    paginator = Paginator(tab_list, items)
    
    tab_items = paginator.page(page)

    pre_column = tab.pre_columns
    richtext_fields = tab.richtext_fields if tab.richtext_fields else []

    total_items = tab_model.objects.filter(**snapshot_filter).count()
    try:
        total_items = tab_model.objects.filter(parent=None).filter(action_status=ActionStatus.IS_ACTIVE).filter(**snapshot_filter).count()
    except:
        pass

    fields = get_model_fields(f'{tab_model.__name__}', pre_column)
    saved_filters = list(SavedFilter.objects.filter(
        parent=ModelChoices.TAB, tab_id=tab.id
    ).values())
    for filter in saved_filters:
        if 'created_at' in filter:
            filter['created_at'] = filter['created_at'].isoformat()
    saved_filters_json = json.dumps(saved_filters)

    label_to_value = {label: value for value, label in ModelChoices.choices}
    s_model_choice = label_to_value.get(tab_model.__name__)

    context = {
        'tab': tab,
        'segment': tab.base_view,
        'parent': tab.sidebar_parent,
        'db_field_names': ordered_fields,
        'field_names': field_names,
        'tab_items': tab_items,
        'pre_column': pre_column,
        'richtext_fields': richtext_fields,
        'total_items': total_items,
        'date_picker_fields': tab_model.date_fields_to_convert if tab_model.date_fields_to_convert else [],
        'int_fields': tab_model.integer_fields if tab_model.integer_fields else [],
        'float_fields': tab_model.float_fields if tab_model.float_fields else [],
        'snapshots': snapshots,
        'selected_snapshot': summary,
        'latest_snapshot': latest_snapshot,
        'is_dynamic_query': tab.is_dynamic_query,
        'query_snapshot': query_snapshot if tab.is_dynamic_query else '',
        'fields': fields,
        'saved_filters': saved_filters_json,
        'saved_filters_len': len(saved_filters),
        'tabs': Tab.objects.filter(base_view=tab.base_view).order_by('created_at').order_by('created_at'),
        'unique_filter': list(set(user_unique_filter + date_unique_filter + int_unique_filter + float_unique_filter)),
        'table_name': tab.content_type.model,
        'selected_rows': selected_rows,
        'selected_rows_qs': tab_list.filter(pk__in=selected_rows),
        'model_name': tab_model.__name__,
        's_model_choice': s_model_choice,
        'tab_notes': tab_notes,
        'notes_form': TabNotesForm(initial={'notes': tab_notes.note})
    }
    return render(request, 'apps/tab/tab_details.html', context)

def remove_tab(request, id):
    tab = Tab.objects.get(id=id)
    SavedFilter.objects.filter(tab_id=id).delete()
    tab.delete()
    return redirect(request.META.get('HTTP_REFERER')) 

def add_tab_description(request, tab_id):
    tab = Tab.objects.get(id=tab_id)
    if request.method == 'POST':
        description = request.POST.get('description')
        tab.description = description
        tab.updated_at = timezone.now()
        tab.save()
        
    return redirect(request.META.get('HTTP_REFERER')) 


def export_tab_csv_view(request, id):
    tab = Tab.objects.get(id=id)
    content_type = ContentType.objects.get(app_label=tab.content_type.app_label, model=tab.content_type.model)
    db_field_names = [field.name for field in content_type.model_class()._meta.get_fields() if not field.is_relation]
    tab_model = content_type.model_class()

    try:
        user_order = ColumnOrder.objects.get(user=request.user, table_name=f'{tab_model.__name__}', tab_id=f"{tab.id}")
        column_names = [col['key'] for col in user_order.column_order if col['key'] is not None]
        ordered_fields = column_names
    except ColumnOrder.DoesNotExist:
        ordered_fields = db_field_names

    selected_rows = SelectedRows.objects.filter(
        model=f'{tab.content_type.model}', 
        model_choice=ModelChoices.TAB,
        tab_id=tab.id
    ).values_list('rows', flat=True)
    selected_rows = [int(item) for row in selected_rows for item in row.split(',')]

    # Snapshot
    latest_snapshot = ''
    summary = None
    snapshot_filter = {}
    try:
        content_type = ContentType.objects.get(model=tab_model.__name__.lower())
        snapshots = InstantUpload.objects.filter(content_type=content_type).order_by('-created_at')
        snapshot = request.GET.get('snapshot')

        if snapshot and not snapshot == 'all':
            summary = InstantUpload.objects.get(id=snapshot)
            snapshot_filter['loader_instance'] = summary.pk
        
        elif snapshot and snapshot == 'all':
            snapshot_filter= {}
        
        elif tab.snapshot == 'latest':
            latest_snapshot = snapshots.latest('created_at')
            snapshot_filter['loader_instance'] = latest_snapshot.pk
        
        elif tab.snapshot and not tab.snapshot == 'all':
            summary = InstantUpload.objects.get(id=int(tab.snapshot))
            snapshot_filter['loader_instance'] = int(tab.snapshot)
        
        elif tab.snapshot and tab.snapshot == 'all':
            snapshot_filter= {}
        

    except:
        pass

    fields = []
    pre_column = tab.pre_columns

    richtext_fields = tab.richtext_fields if tab.richtext_fields else []
    # Show/Hide Filter Fields
    show_fields = tab.hide_show_filters.all().filter(value=False)
    show_fields_keys = {field.key for field in show_fields}

    # Construct final fields list respecting original model order
    for field in ordered_fields:
        # Ensure pre_column fields are not included
        if field not in pre_column and field in show_fields_keys:
            fields.append(field)

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="{tab_model._meta.model_name}.csv"'

    
    filter_string = {}
    combined_q_objects, user_unique_filter, user_query_conditions, user_count_filters = same_key_filter(tab.saved_filters.filter(field_type=FieldType.TEXT), return_count_filters=True)

    pre_filters = {}
    if tab.pre_filters:
        pre_filters = eval(tab.pre_filters)

    # for date range
    date_string, date_unique_filter, date_query_conditions, date_count_filters = common_date_filter(tab.saved_filters.filter(field_type=FieldType.DATE), return_count_filters=True)
    filter_string.update(date_string)
    
    # for integer range
    int_string, int_unique_filter, int_query_conditions, int_count_filters = common_integer_filter(tab.saved_filters.filter(field_type=FieldType.INT), return_count_filters=True)
    filter_string.update(int_string)

    # for float range
    float_string, float_unique_filter, float_query_conditions, float_count_filters = common_float_filter(tab.saved_filters.filter(field_type=FieldType.FLOAT), return_count_filters=True)
    filter_string.update(float_string)

    base_queryset = tab_model.objects.filter(**filter_string).filter(**snapshot_filter).filter(combined_q_objects).filter(**pre_filters)
    order_by = request.GET.get('order_by', 'pk')
    queryset = base_queryset

    if user_query_conditions:
        queryset = queryset.filter(user_query_conditions)
    if date_query_conditions:
        queryset = queryset.filter(date_query_conditions)
    if int_query_conditions:
        queryset = queryset.filter(int_query_conditions)
    if float_query_conditions:
        queryset = queryset.filter(float_query_conditions)

    if user_count_filters:
        queryset = common_count_filter(user_count_filters, base_queryset, queryset, fields)
    elif date_count_filters:
        queryset = common_count_filter(date_count_filters, base_queryset, queryset, fields)
    elif int_count_filters:
        queryset = common_count_filter(int_count_filters, base_queryset, queryset, fields)
    elif float_count_filters:
        queryset = common_count_filter(float_count_filters, base_queryset, queryset, fields)
    else:
        if order_by in ['count', '-count']:
            order_by = 'pk'
        if 'count' in fields:
            fields.remove('count')

    queryset = queryset.order_by(order_by)

    table_name = f"{tab.content_type.app_label}_{tab.content_type.model}"
    if user_unique_filter:
        queryset = common_unique_filter(request, user_unique_filter, queryset, snapshot_filter, table_name)
    if date_unique_filter:
        queryset = common_unique_filter(request, date_unique_filter, queryset, snapshot_filter, table_name)
    if int_unique_filter:
        queryset = common_unique_filter(request, int_unique_filter, queryset, snapshot_filter, table_name)
    if float_unique_filter:
        queryset = common_unique_filter(request, float_unique_filter, queryset, snapshot_filter, table_name)

    tab_list = software_filter(request, queryset, ordered_fields)

    if 'selected_rows' in request.GET:
        tab_list = tab_list.filter(pk__in=selected_rows)

    writer = csv.writer(response)
    writer.writerow(fields) 

    for fav in tab_list:
        row_data = []
        for field in fields:
            if field in tab_model.date_fields_to_convert:
                unix_timestamp = getattr(fav, field)
                if unix_timestamp:
                    date_time = datetime.utcfromtimestamp(unix_timestamp)
                    formatted_date = date_time.strftime('%b %d %Y %I:%M%p')
                    row_data.append(formatted_date)
                else:
                    row_data.append("")

            elif field in richtext_fields:
                attribute = getattr(fav, field, None)
                if attribute and hasattr(attribute, 'html'):
                    row_data.append(attribute.html)
                else:
                    row_data.append("")
            else:
                row_data.append(getattr(fav, field))
        writer.writerow(row_data)

    return response


class ExportTabExcelView(View):
    def get_excel_workbook(self, request, id):
        """Returns just the workbook for easier integration"""
        tab = Tab.objects.get(id=id)
        content_type = ContentType.objects.get(app_label=tab.content_type.app_label, model=tab.content_type.model)
        tab_model = content_type.model_class()
        db_field_names = [field.name for field in tab_model._meta.get_fields() if not field.is_relation]

        try:
            user_order = ColumnOrder.objects.get(user=request.user, table_name=f'{tab_model.__name__}', tab_id=f"{tab.id}")
            column_names = [col['key'] for col in user_order.column_order if col['key'] is not None]
            ordered_fields = column_names
        except ColumnOrder.DoesNotExist:
            ordered_fields = db_field_names

        # Selected rows handling
        selected_rows = SelectedRows.objects.filter(
            model=tab.content_type.model, 
            model_choice=ModelChoices.TAB,
            tab_id=tab.id
        ).values_list('rows', flat=True)
        selected_rows = [int(item) for row in selected_rows for item in row.split(',')]

        # Snapshot handling
        latest_snapshot = ''
        summary = None
        snapshot_filter = {}
        try:
            snapshots = InstantUpload.objects.filter(content_type=content_type).order_by('-created_at')
            snapshot = request.GET.get('snapshot')

            if snapshot and snapshot != 'all':
                summary = InstantUpload.objects.get(id=snapshot)
                snapshot_filter['loader_instance'] = summary.pk
            elif snapshot == 'all':
                snapshot_filter = {}
            elif tab.snapshot == 'latest':
                latest_snapshot = snapshots.latest('created_at')
                snapshot_filter['loader_instance'] = latest_snapshot.pk
            elif tab.snapshot and tab.snapshot != 'all':
                summary = InstantUpload.objects.get(id=int(tab.snapshot))
                snapshot_filter['loader_instance'] = int(tab.snapshot)
            elif tab.snapshot == 'all':
                snapshot_filter = {}
        except:
            pass

        # Field filtering
        fields = []
        pre_column = tab.pre_columns
        richtext_fields = tab.richtext_fields or []
        show_fields = tab.hide_show_filters.filter(value=False)
        show_fields_keys = {field.key for field in show_fields}

        for field in ordered_fields:
            if field not in pre_column and field in show_fields_keys:
                fields.append(field)

        # Filter handling
        filter_string = {}
        combined_q_objects, user_unique_filter, user_query_conditions, user_count_filters = same_key_filter(
            tab.saved_filters.filter(field_type=FieldType.TEXT), return_count_filters=True)

        pre_filters = eval(tab.pre_filters) if tab.pre_filters else {}

        date_string, date_unique_filter, date_query_conditions, date_count_filters = common_date_filter(
            tab.saved_filters.filter(field_type=FieldType.DATE), return_count_filters=True)
        filter_string.update(date_string)

        int_string, int_unique_filter, int_query_conditions, int_count_filters = common_integer_filter(
            tab.saved_filters.filter(field_type=FieldType.INT), return_count_filters=True)
        filter_string.update(int_string)

        float_string, float_unique_filter, float_query_conditions, float_count_filters = common_float_filter(
            tab.saved_filters.filter(field_type=FieldType.FLOAT), return_count_filters=True)
        filter_string.update(float_string)

        # Query construction
        base_queryset = tab_model.objects.filter(**filter_string).filter(**snapshot_filter)\
                                         .filter(combined_q_objects).filter(**pre_filters)
        order_by = request.GET.get('order_by', 'pk')
        queryset = base_queryset

        if user_query_conditions:
            queryset = queryset.filter(user_query_conditions)
        if date_query_conditions:
            queryset = queryset.filter(date_query_conditions)
        if int_query_conditions:
            queryset = queryset.filter(int_query_conditions)
        if float_query_conditions:
            queryset = queryset.filter(float_query_conditions)

        if user_count_filters:
            queryset = common_count_filter(user_count_filters, base_queryset, queryset, fields)
        elif date_count_filters:
            queryset = common_count_filter(date_count_filters, base_queryset, queryset, fields)
        elif int_count_filters:
            queryset = common_count_filter(int_count_filters, base_queryset, queryset, fields)
        elif float_count_filters:
            queryset = common_count_filter(float_count_filters, base_queryset, queryset, fields)
        else:
            if order_by in ['count', '-count']:
                order_by = 'pk'
            if 'count' in fields:
                fields.remove('count')

        queryset = queryset.order_by(order_by)

        # Unique filters
        table_name = f"{tab.content_type.app_label}_{tab.content_type.model}"
        if user_unique_filter:
            queryset = common_unique_filter(request, user_unique_filter, queryset, snapshot_filter, table_name)
        if date_unique_filter:
            queryset = common_unique_filter(request, date_unique_filter, queryset, snapshot_filter, table_name)
        if int_unique_filter:
            queryset = common_unique_filter(request, int_unique_filter, queryset, snapshot_filter, table_name)
        if float_unique_filter:
            queryset = common_unique_filter(request, float_unique_filter, queryset, snapshot_filter, table_name)

        tab_list = software_filter(request, queryset, ordered_fields)

        if 'selected_rows' in request.GET:
            tab_list = tab_list.filter(pk__in=selected_rows)

        # Workbook creation
        wb = Workbook()
        ws = wb.active
        ws.append(fields)

        for fav in tab_list:
            row_data = []
            for field in fields:
                if field in tab_model.date_fields_to_convert:
                    unix_timestamp = getattr(fav, field)
                    if unix_timestamp:
                        epoch = datetime(1970, 1, 1, tzinfo=timezone.utc)
                        date_time = epoch + timedelta(seconds=unix_timestamp)
                        formatted_date = date_time.strftime('%b %d %Y %I:%M%p')
                        row_data.append(formatted_date)
                    else:
                        row_data.append("")
                elif field in richtext_fields:
                    value = getattr(fav, field, None)
                    row_data.append(value.html if hasattr(value, 'html') else "")
                else:
                    row_data.append(getattr(fav, field))
            ws.append(row_data)

        return wb

    def get(self, request, id):
        """Original endpoint that returns HTTP response"""
        wb = self.get_excel_workbook(request, id)
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        tab = Tab.objects.get(id=id)
        content_type = ContentType.objects.get(app_label=tab.content_type.app_label, model=tab.content_type.model)
        tab_model = content_type.model_class()
        response['Content-Disposition'] = f'attachment; filename="{tab_model._meta.model_name}.xlsx"'
        wb.save(response)
        return response


def add_tab_notes(request, tab_id):
    tab = Tab.objects.get(id=tab_id)
    tab_notes, _ = TabNotes.objects.get_or_create(tab=tab)
    if request.method == "POST":
        notes = request.POST.get('notes')
        tab_notes.note = notes
        tab_notes.save()

        return redirect(request.META.get('HTTP_REFERER'))
    
    return redirect(request.META.get('HTTP_REFERER'))


def rename_tab(request, tab_id):
    tab = Tab.objects.get(id=tab_id)
    if request.method == 'POST':
        name = request.POST.get('name')
        tab.name = name
        tab.save()
        
        return redirect(request.META.get('HTTP_REFERER'))
    
    return redirect(request.META.get('HTTP_REFERER'))
